"""Relatório mensal: tabelas, variação contra o mês anterior, Excel e HTML do PDF."""

from datetime import date

import pytest
from openpyxl import load_workbook

from monitor.relatorio.dados import Coluna, montar_relatorio, nome_mes, variacao
from monitor.relatorio.excel import exportar_excel
from monitor.relatorio.pdf import montar_html


def _dados() -> dict:
    instituicoes = [[1, "Banco A", "a", None], [2, "Banco B", "b", None]]
    perfis = [
        [10, 1, "app_store", "111", None], [11, 1, "google_play", "com.a", None],
        [20, 2, "app_store", "222", None],
        [30, 1, "consumidor_gov", "300", None],
        [40, 1, "reclame_aqui", "banco-a", None],
        [50, 1, "instagram", "bancoa", None],
    ]
    historico = [
        # setembro
        [10, "nota_media", "2026-09-20", 4.5, None, None, None],
        [11, "nota_media", "2026-09-20", 4.0, None, None, None],
        [30, "pct_resolvidas", "2026-09-20", 80.0, None, "2026-01-01", "2026-09-20"],
        [30, "total_reclamacoes", "2026-09-20", 100, None, "2026-01-01", "2026-09-20"],
        [40, "status_reputacao", "2026-09-20", None, "Bom", "2026-03-01", "2026-08-31"],
        [50, "seguidores", "2026-09-20", 1000, None, None, None],
        [50, "posts_periodo", "2026-09-20", 1, None, "2026-08-01", "2026-08-31"],
        # outubro (duas coletas: vale a última do mês)
        [10, "nota_media", "2026-10-11", 4.6, None, None, None],
        [10, "nota_media", "2026-10-24", 4.7, None, None, None],
        [11, "nota_media", "2026-10-24", 3.9, None, None, None],
        [20, "nota_media", "2026-10-24", 4.1, None, None, None],
        [30, "pct_resolvidas", "2026-10-24", 82.5, None, "2026-01-01", "2026-10-24"],
        [30, "total_reclamacoes", "2026-10-24", 150, None, "2026-01-01", "2026-10-24"],
        [40, "status_reputacao", "2026-10-24", None, "Ótimo", "2026-04-01", "2026-09-30"],
        [50, "seguidores", "2026-10-24", 1100, None, None, None],
        [50, "posts_periodo", "2026-10-24", 2, None, "2026-09-01", "2026-09-30"],
    ]
    posts = [
        [50, "2026-08-10", "2026-08-01", 30, False, 1, "IMAGE", "https://www.instagram.com/p/x/"],
        [50, "2026-09-05", "2026-09-01", 100, False, 2, "IMAGE", "https://www.instagram.com/p/a/"],
        [50, "2026-09-20", "2026-09-01", None, True, 3, "VIDEO", "https://www.instagram.com/p/b/"],
    ]
    metricas = [["nota_media", "Nota média", "nota"], ["seguidores", "Seguidores", "contagem"]]
    return {"instituicoes": instituicoes, "perfis": perfis, "metricas": metricas, "historico": historico,
            "posts": posts, "execucoes": [], "gerado_em": "24/10/2026 10:00"}


def test_valores_do_mes_e_do_mes_anterior():
    rel = montar_relatorio(_dados(), "2026-10", "teste")
    assert (rel.data_ref, rel.data_anterior) == (date(2026, 10, 24), date(2026, 9, 20))
    resumo = {l.rotulo: l for l in rel.resumo.linhas}
    a = resumo["Banco A"]
    assert a.valores[0] == 4.7 and a.anteriores[0] == 4.5  # última coleta do mês, não a de 11/10
    assert a.valores[3] == 82.5 and a.anteriores[3] == 80.0
    assert a.valores[6] == "Ótimo" and a.anteriores[6] == "Bom"
    assert resumo["Banco B"].anteriores[0] is None  # sem dado no mês anterior: nunca zero


def test_instagram_usa_a_janela_de_cada_levantamento():
    rel = montar_relatorio(_dados(), "2026-10")
    (linha,) = next(t for t in rel.tabelas if t.id == "instagram").linhas
    # Outubro: posts de setembro (2, um com curtidas ocultas); setembro: posts de agosto.
    assert linha.valores == [1100, 2, 100, 100]
    assert linha.anteriores == [1000, 1, 30, 30]
    assert [p["curtidas"] for p in rel.top_posts] == [100, None]


def test_mes_sem_levantamento_anterior():
    rel = montar_relatorio(_dados(), "2026-09")
    assert rel.data_anterior is None
    assert "Ainda não há levantamento do mês anterior" in montar_html(rel)


def test_historico_guarda_todos_os_levantamentos():
    rel = montar_relatorio(_dados(), "2026-10")
    # Todas as coletas ficam, inclusive as duas de outubro do mesmo app.
    assert len(rel.todos_levantamentos) == len(_dados()["historico"])
    # Linha: mês, instituição, fonte, perfil, métrica, código, valor, ...
    mensal = {(l[0], l[1], l[2], l[5]): l for l in rel.historico_mensal}
    assert mensal[("2026-10", "Banco A", "App Store", "nota_media")][6] == 4.7  # último de outubro


@pytest.mark.parametrize(
    ("atual", "anterior", "coluna", "esperado"),
    [
        (82.5, 80.0, Coluna("x", "pct"), ("▲ 2,5 p.p.", "melhor")),
        (5.0, 6.0, Coluna("x", "dias", "menor"), ("▼ 1,0 d", "melhor")),
        (150, 100, Coluna("x", "contagem", "volume"), ("▲ 50", "neutro")),
        (4.501, 4.5, Coluna("x", "nota"), ("=", "neutro")),
        (4.5, None, Coluna("x", "nota"), None),
    ],
)
def test_variacao(atual, anterior, coluna, esperado):
    assert variacao(atual, anterior, coluna) == esperado


def test_excel_tem_abas_formulas_e_graficos(tmp_path):
    rel = montar_relatorio(_dados(), "2026-10", "teste")
    caminho = exportar_excel(rel, tmp_path / "r.xlsx")
    wb = load_workbook(caminho)
    assert wb.sheetnames == ["Resumo", "Lojas", "Lojas por app", "Consumidor.gov", "Reclame Aqui", "Instagram",
                             "Posts Instagram", "Histórico mensal", "Todos os levantamentos", "Sobre"]
    ws = wb["Consumidor.gov"]
    valores = [c.value for linha in ws.iter_rows() for c in linha]
    assert 0.825 in valores and 0.8 in valores  # percentual guardado como fração
    assert any(isinstance(v, str) and v.startswith("=IF(OR(") and "*100" in v for v in valores)
    assert len(wb["Lojas"]._charts) == 1 and len(wb["Instagram"]._charts) == 1


def test_lojas_em_tabela_com_media_e_diferenca():
    rel = montar_relatorio(_dados(), "2026-10")
    lojas = next(t for t in rel.tabelas if t.id == "lojas")
    assert [c.titulo for c in lojas.colunas] == ["App Store", "Google Play", "Média das lojas", "Diferença (AS − GP)"]
    a = next(l for l in lojas.linhas if l.rotulo == "Banco A")
    assert a.valores[:3] == [4.7, 3.9, pytest.approx(4.3)] and a.valores[3] == pytest.approx(0.8)
    assert lojas.linhas[0].rotulo == "Banco A"  # ordenado pela média das lojas
    assert "<svg" not in montar_html(rel)


def test_html_do_pdf_escapa_textos_e_mostra_variacao():
    dados = _dados()
    dados["instituicoes"][0][1] = "Banco <A>"
    html = montar_html(montar_relatorio(dados, "2026-10"))
    assert "Banco &lt;A&gt;" in html and "Banco <A>" not in html
    assert "▲ 2,5 p.p." in html
    assert nome_mes("2026-10") == "outubro de 2026"
