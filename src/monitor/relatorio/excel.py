"""Exporta o relatório para Excel (.xlsx) com formatação, variações e gráficos nativos."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from monitor.relatorio.dados import Coluna, Relatorio, Tabela, fmt_data

FONTE = "Arial"
COR_TINTA = "1D2433"
COR_SUAVE = "5B6475"
COR_CABECALHO = "EEF2F8"
COR_GRUPO = "DCE6F4"
COR_LINHA = "E3E7ED"
COR_MELHOR = "0A7D0A"
COR_PIOR = "C13030"
SERIE_1, SERIE_2 = "2A78D6", "EB6834"

FORMATOS = {
    "nota": ("0.00", '+0.00;-0.00;0.00'),
    "pct": ("0.0%", '+0.0" p.p.";-0.0" p.p.";0.0" p.p."'),
    "contagem": ("#,##0", "+#,##0;-#,##0;0"),
    "dias": ('0.0" d"', '+0.0" d";-0.0" d";0.0" d"'),
    "decimal": ("#,##0.0", "+#,##0.0;-#,##0.0;0.0"),
}

_fino = Side(style="thin", color=COR_LINHA)
BORDA = Border(bottom=_fino)


def _fonte(**kw) -> Font:
    return Font(name=FONTE, size=kw.pop("size", 10), color=kw.pop("color", COR_TINTA), **kw)


def _valor_celula(valor, coluna: Coluna):
    if valor is None:
        return None
    if coluna.formato == "pct":
        return valor / 100  # Excel guarda percentual como fração
    return valor


def _cabecalho_folha(ws: Worksheet, titulo: str, rel: Relatorio, tabela: Tabela | None = None) -> int:
    ws["A1"] = titulo
    ws["A1"].font = _fonte(size=15, bold=True)
    comparacao = (f"comparado com o levantamento de {fmt_data(rel.data_anterior)}" if rel.data_anterior
                  else "sem levantamento do mês anterior para comparar")
    ws["A2"] = f"Referência: {rel.titulo_mes} · levantamento de {fmt_data(rel.data_ref)} · {comparacao}"
    ws["A2"].font = _fonte(color=COR_SUAVE)
    linha = 3
    if tabela and tabela.janela:
        ws.cell(linha, 1, f"Atual: {tabela.janela}").font = _fonte(color=COR_SUAVE, size=9)
        linha += 1
    if tabela and tabela.janela_anterior and rel.data_anterior:
        ws.cell(linha, 1, f"Mês anterior: {tabela.janela_anterior}").font = _fonte(color=COR_SUAVE, size=9)
        linha += 1
    return linha + 1


def escrever_tabela(ws: Worksheet, tabela: Tabela, rel: Relatorio, inicio: int) -> dict:
    """Escreve a tabela a partir da linha `inicio`. Devolve as posições usadas (para gráficos)."""
    com_anterior = rel.data_anterior is not None
    grupo_linha, sub_linha = inicio, inicio + 1
    ws.cell(grupo_linha, 1, tabela.rotulo_linha)
    ws.merge_cells(start_row=grupo_linha, start_column=1, end_row=sub_linha, end_column=1)
    posicoes = {}
    col = 2
    for k, coluna in enumerate(tabela.colunas):
        partes = ["Atual"]
        if com_anterior:
            partes += ["Mês anterior"] + (["Variação"] if coluna.formato != "texto" else [])
        titulo = f"{coluna.grupo} · {coluna.titulo}" if coluna.grupo else coluna.titulo
        ws.cell(grupo_linha, col, titulo)
        if len(partes) > 1:
            ws.merge_cells(start_row=grupo_linha, start_column=col, end_row=grupo_linha, end_column=col + len(partes) - 1)
        for j, parte in enumerate(partes):
            ws.cell(sub_linha, col + j, parte)
        posicoes[k] = (col, len(partes))
        col += len(partes)
    ultima_col = col - 1

    for c in range(1, ultima_col + 1):
        for r, cor in ((grupo_linha, COR_GRUPO), (sub_linha, COR_CABECALHO)):
            cel = ws.cell(r, c)
            cel.font = _fonte(bold=True, size=9 if r == sub_linha else 10)
            cel.fill = PatternFill("solid", fgColor=cor)
            cel.alignment = Alignment(horizontal="center" if c > 1 else "left", vertical="center", wrap_text=True)
            cel.border = BORDA
    ws.row_dimensions[grupo_linha].height = 30

    primeira = sub_linha + 1
    for i, linha in enumerate(tabela.linhas):
        r = primeira + i
        nome = ws.cell(r, 1, linha.rotulo)
        nome.font = _fonte()
        nome.border = BORDA
        for k, coluna in enumerate(tabela.colunas):
            c, largura = posicoes[k]
            formato, formato_var = FORMATOS.get(coluna.formato, ("@", "@"))
            atual = ws.cell(r, c, _valor_celula(linha.valores[k], coluna))
            atual.number_format = formato
            atual.font = _fonte(bold=True)
            celulas = [atual]
            if com_anterior:
                anterior = ws.cell(r, c + 1, _valor_celula(linha.anteriores[k], coluna))
                anterior.number_format = formato
                anterior.font = _fonte(color=COR_SUAVE)
                celulas.append(anterior)
                if largura == 3:
                    a, b = f"{get_column_letter(c)}{r}", f"{get_column_letter(c + 1)}{r}"
                    fator = "*100" if coluna.formato == "pct" else ""
                    var = ws.cell(r, c + 2, f'=IF(OR({a}="",{b}=""),"",({a}-{b}){fator})')
                    var.number_format = formato_var
                    var.font = _fonte(size=9, color=COR_SUAVE)
                    celulas.append(var)
            for cel in celulas:
                cel.border = BORDA
                cel.alignment = Alignment(horizontal="right" if coluna.formato != "texto" else "center")
    ultima = primeira + len(tabela.linhas) - 1

    # Variação colorida: verde quando melhora, vermelho quando piora (o sinal também aparece no número).
    if com_anterior and tabela.linhas:
        for k, coluna in enumerate(tabela.colunas):
            c, largura = posicoes[k]
            if largura != 3 or coluna.sentido == "volume":
                continue
            letra = get_column_letter(c + 2)
            faixa = f"{letra}{primeira}:{letra}{ultima}"
            sobe, desce = (COR_MELHOR, COR_PIOR) if coluna.sentido == "maior" else (COR_PIOR, COR_MELHOR)
            celula = f"{letra}{primeira}"
            ws.conditional_formatting.add(faixa, FormulaRule(
                formula=[f"AND(ISNUMBER({celula}),{celula}>0)"], font=Font(name=FONTE, size=9, color=sobe, bold=True)))
            ws.conditional_formatting.add(faixa, FormulaRule(
                formula=[f"AND(ISNUMBER({celula}),{celula}<0)"], font=Font(name=FONTE, size=9, color=desce, bold=True)))

    ws.column_dimensions["A"].width = 38
    for c in range(2, ultima_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 14
    if tabela.nota:
        ws.cell(ultima + 2, 1, tabela.nota).font = _fonte(size=9, color=COR_SUAVE, italic=True)
    return {"primeira": primeira, "ultima": ultima, "posicoes": posicoes, "sub": sub_linha, "ultima_col": ultima_col}


def _grafico_barras(ws: Worksheet, titulo: str, pos: dict, colunas: list[tuple[int, str, str]],
                    ancora: str, eixo: tuple[float, float] | None = None) -> None:
    grafico = BarChart()
    grafico.type = "bar"
    grafico.style = 10
    grafico.title = titulo
    grafico.y_axis.majorGridlines = None
    grafico.gapWidth = 60
    if eixo:
        grafico.y_axis.scaling.min, grafico.y_axis.scaling.max = eixo
    categorias = Reference(ws, min_col=1, min_row=pos["primeira"], max_row=pos["ultima"])
    for coluna, nome, cor in colunas:
        dados = Reference(ws, min_col=coluna, min_row=pos["primeira"], max_row=pos["ultima"])
        grafico.add_data(dados, titles_from_data=False)
        serie = grafico.series[-1]
        serie.tx = SeriesLabel(v=nome)
        serie.graphicalProperties.solidFill = cor
        serie.graphicalProperties.line.solidFill = cor
    grafico.set_categories(categorias)
    grafico.x_axis.scaling.orientation = "maxMin"  # mesma ordem da tabela
    grafico.legend.position = "t"
    grafico.height = max(8, 0.55 * (pos["ultima"] - pos["primeira"] + 1))
    grafico.width = 16
    ws.add_chart(grafico, ancora)


def _lista(ws: Worksheet, inicio: int, cabecalho: list[str], linhas: list[list], formatos: dict[int, str]) -> None:
    for c, titulo in enumerate(cabecalho, start=1):
        cel = ws.cell(inicio, c, titulo)
        cel.font = _fonte(bold=True, size=9)
        cel.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        cel.border = BORDA
    for i, linha in enumerate(linhas, start=inicio + 1):
        for c, valor in enumerate(linha, start=1):
            cel = ws.cell(i, c, valor)
            cel.font = _fonte()
            if c in formatos:
                cel.number_format = formatos[c]
    ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(cabecalho))}{max(inicio, inicio + len(linhas))}"
    ws.freeze_panes = ws.cell(inicio + 1, 1)


def exportar_excel(rel: Relatorio, caminho: Path) -> Path:
    wb = Workbook()

    # Resumo
    ws = wb.active
    ws.title = "Resumo"
    linha = _cabecalho_folha(ws, f"Relatório de reputação digital · {rel.titulo_mes}", rel)
    for k, (rotulo, valor, detalhe) in enumerate(rel.kpis):
        ws.cell(linha + k, 1, rotulo).font = _fonte(color=COR_SUAVE)
        ws.cell(linha + k, 2, valor).font = _fonte(bold=True, size=12)
        ws.cell(linha + k, 3, detalhe).font = _fonte(size=9, color=COR_SUAVE)
    pos = escrever_tabela(ws, rel.resumo, rel, linha + len(rel.kpis) + 1)
    ws.freeze_panes = ws.cell(pos["primeira"], 2)

    nomes = {"lojas": "Lojas", "lojas_apps": "Lojas por app", "consumidor_gov": "Consumidor.gov",
             "reclame_aqui": "Reclame Aqui", "instagram": "Instagram"}
    for tabela in rel.tabelas:
        ws = wb.create_sheet(nomes.get(tabela.id, tabela.id)[:31])
        inicio = _cabecalho_folha(ws, f"{tabela.titulo} · {rel.titulo_mes}", rel, tabela)
        pos = escrever_tabela(ws, tabela, rel, inicio)
        ws.freeze_panes = ws.cell(pos["primeira"], 2)
        ancora = f"{get_column_letter(pos['ultima_col'] + 2)}{inicio}"
        if tabela.id == "lojas" and tabela.linhas:
            _grafico_barras(ws, "Nota média por loja", pos,
                            [(pos["posicoes"][0][0], "App Store", SERIE_1), (pos["posicoes"][1][0], "Google Play", SERIE_2)],
                            ancora, eixo=(3, 5))
        if tabela.id == "instagram" and tabela.linhas:
            _grafico_barras(ws, "Seguidores", pos, [(pos["posicoes"][0][0], "Seguidores", SERIE_1)], ancora)

    # Posts do Instagram
    ws = wb.create_sheet("Posts Instagram")
    inicio = _cabecalho_folha(ws, f"Posts do Instagram · {rel.titulo_mes}", rel)
    tipos = {"IMAGE": "imagem", "VIDEO": "vídeo", "CAROUSEL_ALBUM": "carrossel"}
    _lista(ws, inicio, ["Perfil", "Data", "Tipo", "Curtidas", "Comentários", "Link"],
           [[p["perfil"], p["data"], tipos.get(p["tipo"], p["tipo"]),
             p["curtidas"] if p["curtidas"] is not None else ("ocultas" if p["ocultas"] else None),
             p["comentarios"], p["link"]] for p in rel.top_posts],
           {2: "DD/MM/YYYY", 4: "#,##0", 5: "#,##0"})
    for r in range(inicio + 1, inicio + 1 + len(rel.top_posts)):
        link = ws.cell(r, 6)
        if isinstance(link.value, str) and link.value.startswith("https://www.instagram.com/"):
            link.hyperlink = link.value
            link.value = "abrir"
            link.font = _fonte(color=SERIE_1, underline="single")
    for letra, largura in zip("ABCDEF", (36, 12, 12, 12, 13, 10), strict=True):
        ws.column_dimensions[letra].width = largura

    # Histórico mensal e todos os levantamentos (formato longo, com filtros)
    cab_hist = ["Instituição", "Fonte", "Perfil", "Métrica", "Código", "Valor", "Texto", "Janela início", "Janela fim"]
    ws = wb.create_sheet("Histórico mensal")
    inicio = _cabecalho_folha(ws, "Histórico mensal (último levantamento de cada mês)", rel)
    _lista(ws, inicio, ["Mês", *cab_hist, "Data da coleta"], rel.historico_mensal,
           {7: "#,##0.00", 9: "DD/MM/YYYY", 10: "DD/MM/YYYY", 11: "DD/MM/YYYY"})
    for letra, largura in zip("ABCDEFGHIJK", (9, 24, 16, 22, 30, 20, 13, 14, 12, 12, 13), strict=True):
        ws.column_dimensions[letra].width = largura

    ws = wb.create_sheet("Todos os levantamentos")
    inicio = _cabecalho_folha(ws, "Todos os levantamentos", rel)
    _lista(ws, inicio, ["Data da coleta", *cab_hist], rel.todos_levantamentos,
           {1: "DD/MM/YYYY", 7: "#,##0.00", 9: "DD/MM/YYYY", 10: "DD/MM/YYYY"})
    for letra, largura in zip("ABCDEFGHIJ", (13, 24, 16, 22, 30, 20, 13, 14, 12, 12), strict=True):
        ws.column_dimensions[letra].width = largura

    # Sobre
    ws = wb.create_sheet("Sobre")
    textos = [
        ("Relatório de reputação digital", True),
        (f"Mês de referência: {rel.titulo_mes}. Valores do último levantamento até o fim do mês ({fmt_data(rel.data_ref)}).", False),
        ("Mês anterior: último levantamento até o fim do mês anterior"
         + (f" ({fmt_data(rel.data_anterior)})." if rel.data_anterior else ": não há levantamento, por isso não há variação."), False),
        ("Variação: atual menos mês anterior. Verde quando melhora, vermelho quando piora; volumes (reclamações, posts) ficam sem cor.", False),
        ("Fontes e janelas: lojas = nota atual; Consumidor.gov.br = ano corrente; Reclame Aqui = últimos 6 meses; "
         "Instagram = seguidores atuais e posts do último mês inteiro.", False),
        ("Nas abas de histórico, percentuais estão na escala de 0 a 100.", False),
        (f"Gerado em {rel.gerado_em} a partir de {rel.origem}.", False),
    ]
    for i, (texto, negrito) in enumerate(textos, start=1):
        cel = ws.cell(i, 1, texto)
        cel.font = _fonte(bold=negrito, size=14 if negrito else 10)
        cel.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    for folha in wb.worksheets:
        folha.sheet_view.showGridLines = False
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho
